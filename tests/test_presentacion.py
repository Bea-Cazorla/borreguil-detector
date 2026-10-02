"""Pruebas de cómo se MUESTRAN los datos: 4 decimales y nombres largos sin cortar.

Se ejecutan con `python tests/test_presentacion.py`. Lo que se comprueba es lo que
se ve (globos del mapa, Excel); los valores guardados conservan toda su precisión.
"""
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import borreguil_pipeline as bp   # noqa: E402

bp._imports()

ID_LARGO = 'all_borreguil_lagunas500ptos_epsg25830.21'


def punto(**extra):
    r = {'ID': ID_LARGO, 'lon': -3.3312, 'lat': 37.0493123456,
         'decision': 'BORREGUIL PROBABLE', 'rf_proba': 0.9333333,
         'mat_signature': 'BORREGUIL_LIKE',
         'elev_dem_m': 2935.835205078125, 'slope_deg': 14.577339172363281,
         'ndvi_late': 0.2185155749320984, 'clre_late': 0.28997838497161865,
         'ambiente': 'ladera', 'humedad': 'seco', 'pureza': 'mixto-agua',
         'tipo_borreguil': 'ladera · seco · mixto-agua'}
    r.update(extra)
    return r


def test_numeros_con_cuatro_decimales():
    assert bp.fmt_num(2935.835205078125) == '2935.8352'
    assert bp.fmt_num(0.28997838497161865) == '0.2900'
    assert bp.fmt_num('0.2185155749320984') == '0.2185'     # en un KML llegan como texto
    assert bp.fmt_num(3) == '3.0000'
    for vacio in (None, float('nan'), '', 'abc'):
        assert bp.fmt_num(vacio) == '—'


def test_ids_largos_abreviados_conservando_lo_que_los_distingue():
    ids = [f'all_borreguil_lagunas500ptos_epsg25830.{n}' for n in (1, 21, 41, 441)]
    cortos = bp.abreviar_ids(ids)
    assert cortos == ['…epsg25830.1', '…epsg25830.21', '…epsg25830.41', '…epsg25830.441']
    assert len(set(cortos)) == len(set(ids))                # siguen distinguiéndose
    # Los cortos no se tocan, tampoco cuando conviven con los largos.
    assert bp.abreviar_ids(['p1', 'campo_22']) == ['p1', 'campo_22']
    assert bp.abreviar_ids(['rand_0001'] + ids)[:2] == ['rand_0001', '…epsg25830.1']
    assert bp.abreviar_ids([7, None]) == ['7', '']
    # Un nombre largo suelto se acorta por el medio, conservando el final.
    assert bp.abreviar_ids(['p1', ids[1]]) == ['p1', 'all_bo…ptos_epsg25830.21']
    # Si al acortar dos nombres distintos quedaran iguales, no se acorta ninguno.
    parecidos = ['A' * 20 + '1' + 'B' * 20, 'A' * 20 + '2' + 'B' * 20]
    assert bp.abreviar_ids(parecidos) == parecidos


def test_globo_del_mapa_con_cuatro_decimales():
    h = bp.popup_html(punto())
    assert 'Altitud: 2935.8352 m' in h and 'Slope: 14.5773°' in h
    assert 'NDVI: 0.2185' in h and 'Clre: 0.2900' in h and 'RF: 93%' in h
    assert '2935.83520' not in h and '0.21851' not in h     # nada de colas largas


def test_nombre_con_puntos_de_corte_en_los_separadores():
    h = bp.nombre_html(ID_LARGO)
    assert h == 'all_<wbr>borreguil_<wbr>lagunas500ptos_<wbr>epsg25830.<wbr>21'
    assert h.replace('<wbr>', '') == ID_LARGO          # el texto es el nombre exacto
    assert bp.nombre_html('a<b>&c') == 'a&lt;b&gt;&amp;c'
    assert bp.nombre_html(17) == '17'


def test_globo_parte_los_nombres_largos():
    h = bp.popup_html(punto())
    assert bp.nombre_html(ID_LARGO) in h
    # Sin esta regla el nombre, que no tiene espacios, se salía del globo.
    assert 'overflow-wrap:anywhere' in h and 'max-width:280px' in h


def test_globo_muestra_el_tipo_y_si_esta_revisado():
    h = bp.popup_html(punto())
    assert 'Tipo: <b>ladera · seco · mixto-agua</b>' in h and 'revisado' not in h
    h = bp.popup_html(punto(tipo_revisado='si'))
    assert '✓ revisado' in h
    assert 'Tipo' not in bp.popup_html(punto(tipo_borreguil='', decision='NO BORREGUIL'))


def test_globo_en_otro_idioma():
    en = {'BORREGUIL PROBABLE': 'LIKELY BORREGUIL',
          'ladera · seco · mixto-agua': 'slope · dry · mixed-water'}
    h = bp.popup_html(punto(tipo_revisado='si'), lambda v: en.get(v, v),
                      {'tipo': 'Type', 'revisado': 'reviewed', 'patron': 'Pattern',
                       'altitud': 'Elevation'})
    assert 'LIKELY BORREGUIL' in h and 'Type: <b>slope · dry · mixed-water</b>' in h
    assert '✓ reviewed' in h and 'Pattern: BORREGUIL_LIKE' in h and 'Elevation: 2935.8352 m' in h


def test_globo_no_se_rompe_con_textos_raros_ni_huecos():
    h = bp.popup_html({'ID': '<b>punto</b> & co', 'lon': 0, 'lat': 0})
    assert '&lt;b&gt;punto&lt;/<wbr>b&gt; &amp; co' in h and '<b>punto</b>' not in h
    assert 'SIN PREDICCIÓN' in h and 'RF: —' in h and 'Altitud: — m' in h
    nan = float('nan')
    h = bp.popup_html(punto(decision=nan, mat_signature=nan, rf_proba=nan, elev_dem_m=nan))
    assert 'SIN PREDICCIÓN' in h and 'Patrón: —' in h and 'nan' not in h.lower()


def test_mapa_html_guardado_usa_el_mismo_globo():
    with tempfile.TemporaryDirectory() as d:
        ruta = os.path.join(d, 'mapa.html')
        bp.save_map([punto(), punto(ID='otro', lon=-3.32, tipo_revisado='si')], ruta)
        html = Path(ruta).read_text(encoding='utf-8')
    assert '2935.8352' in html and '2935.83520' not in html
    assert 'overflow-wrap:anywhere' in html and 'revisado' in html


def test_excel_cuatro_decimales_y_columnas_a_lo_ancho_del_nombre():
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    with tempfile.TemporaryDirectory() as d:
        ruta = os.path.join(d, 'puntos.xlsx')
        bp.save_xlsx([punto(), punto(ID='corto')], ruta)
        ws = load_workbook(ruta).active
        cab = {c.value: i for i, c in enumerate(ws[1], 1)}
        alt = ws.cell(row=2, column=cab['Altitud (m)'])
        assert alt.value == 2935.835205078125           # el dato, entero
        assert alt.number_format == '0.0000'            # lo que se ve, 4 decimales
        assert ws.cell(row=2, column=cab['lat']).number_format == '0.000000'
        ancho_id = ws.column_dimensions[get_column_letter(cab['ID'])].width
        assert ancho_id >= len(ID_LARGO), ancho_id
        ancho_cab = ws.column_dimensions[
            get_column_letter(cab['Ambiente (propuesta app)'])].width
        assert ancho_cab >= len('Ambiente (propuesta app)')


if __name__ == '__main__':
    pruebas = [(n, f) for n, f in sorted(globals().items())
               if n.startswith('test_') and callable(f)]
    fallos = 0
    for nombre, f in pruebas:
        try:
            f()
            print(f'  ok     {nombre}')
        except Exception as e:                      # noqa: BLE001
            fallos += 1
            print(f'  FALLO  {nombre}: {type(e).__name__}: {e}')
    print(f'{len(pruebas) - fallos}/{len(pruebas)} pruebas correctas')
    sys.exit(1 if fallos else 0)
